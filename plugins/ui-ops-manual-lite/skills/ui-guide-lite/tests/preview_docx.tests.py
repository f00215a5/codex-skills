from __future__ import annotations

import json
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path

from docx import Document
from PIL import Image


ROOT = Path(__file__).parents[1]
SCRIPTS = ROOT / "scripts"
BUILD = SCRIPTS / "build_docx.py"
PREVIEW = SCRIPTS / "preview_docx.py"


def run(script: Path, *args: str) -> subprocess.CompletedProcess[str]:
    return subprocess.run(
        [sys.executable, str(script), *args],
        check=False,
        capture_output=True,
        text=True,
        encoding="utf-8",
    )


class PreviewDocxTests(unittest.TestCase):
    def setUp(self) -> None:
        self.temp_dir = tempfile.TemporaryDirectory()
        self.work = Path(self.temp_dir.name)
        (self.work / "annotated").mkdir()
        Image.new("RGB", (200, 100), "white").save(self.work / "annotated" / "step.png")
        self.manifest = self.work / "manual.json"
        self.manifest.write_text(json.dumps({
            "title": "預覽測試",
            "subtitle": "內容讀回",
            "applicableScreens": ["測試畫面"],
            "revision": {"version": "1.0", "date": "2026-01-01", "summary": "測試"},
            "usageReminders": ["測試提醒"],
            "commonRules": ["測試規則"],
            "chapter": {
                "title": "測試功能",
                "entry": {"caption": "功能入口。", "image": "annotated/step.png"},
                "sections": [{
                    "title": "查詢",
                    "steps": [{"caption": "紅框 1：點選查詢。", "image": "annotated/step.png"}],
                    "fields": [{"name": "關鍵字", "definition": "查詢文字。", "required": "選填", "limits": "文字"}],
                    "impact": "無資料變更。",
                    "verification": "結果可見。",
                }],
            },
            "updateLog": [{"version": "1.0", "date": "2026-01-01", "changes": "測試"}],
        }, ensure_ascii=False), encoding="utf-8")
        self.docx = self.work / "manual.docx"
        built = run(BUILD, "--manifest", str(self.manifest), "--output", str(self.docx))
        self.assertEqual(built.returncode, 0, built.stderr)

    def tearDown(self) -> None:
        self.temp_dir.cleanup()

    def test_preview_walks_content_and_scales_inline_images(self) -> None:
        output = self.work / "preview"
        result = run(PREVIEW, "--docx", str(self.docx), "--output-dir", str(output), "--dpi", "96")

        self.assertEqual(result.returncode, 0, result.stderr)
        report = json.loads((output / "preview.json").read_text(encoding="utf-8"))
        self.assertEqual(report["render_visual"]["status"], "not_performed")
        self.assertEqual(report["status"], "pass")
        self.assertGreaterEqual(len(report["images"]), 2)
        self.assertEqual(report["images"][0]["scaled_size_px"], {"width": 576, "height": 288})
        self.assertTrue(any(block["type"] == "table" for block in report["blocks"]))
        captions = [block.get("caption") for block in report["blocks"] if block.get("caption")]
        self.assertIn("紅框 1：點選查詢。", captions)
        html = (output / "preview.html").read_text(encoding="utf-8")
        self.assertIn("render_visual: not_performed", html)
        self.assertIn("內容預覽，非Word分頁/字型渲染", html)
        self.assertIn("images/image-001.png", html)
        for image in report["images"]:
            self.assertRegex(image["source_blob_sha256"], r"^[0-9a-f]{64}$")
            self.assertRegex(image["preview_sha256"], r"^[0-9a-f]{64}$")

    def test_preview_rejects_non_positive_dpi(self) -> None:
        for dpi in ("0", "-1", "601", "nan", "inf"):
            with self.subTest(dpi=dpi):
                result = run(
                    PREVIEW,
                    "--docx",
                    str(self.docx),
                    "--output-dir",
                    str(self.work / f"preview-{dpi}"),
                    "--dpi",
                    dpi,
                )
                self.assertEqual(result.returncode, 2, result.stderr)
                self.assertIn("dpi", result.stderr.lower())

    def test_preview_refuses_to_overwrite_existing_checkpoint(self) -> None:
        output = self.work / "preview"
        output.mkdir()
        (output / "checkpoint.json").write_text("{}", encoding="utf-8")
        result = run(PREVIEW, "--docx", str(self.docx), "--output-dir", str(output))
        self.assertEqual(result.returncode, 2)
        self.assertIn("overwrite", result.stderr.lower())

    def test_empty_document_is_blocked_and_reports_scope(self) -> None:
        empty_docx = self.work / "empty.docx"
        Document().save(empty_docx)
        output = self.work / "empty-preview"
        result = run(PREVIEW, "--docx", str(empty_docx), "--output-dir", str(output))
        self.assertEqual(result.returncode, 1, result.stderr)
        report = json.loads((output / "preview.json").read_text(encoding="utf-8"))
        self.assertEqual(report["status"], "blocked")
        self.assertIn("content_empty", {finding["code"] for finding in report["findings"]})
        self.assertEqual(report["coverage"]["headers_footers"], "not_walked")

    def test_table_cell_image_is_explicitly_blocked(self) -> None:
        table_docx = self.work / "table-image.docx"
        document = Document()
        table = document.add_table(rows=1, cols=1)
        run_obj = table.cell(0, 0).paragraphs[0].add_run()
        run_obj.add_picture(str(self.work / "annotated" / "step.png"))
        document.save(table_docx)
        output = self.work / "table-image-preview"
        result = run(PREVIEW, "--docx", str(table_docx), "--output-dir", str(output))
        self.assertEqual(result.returncode, 1, result.stderr)
        report = json.loads((output / "preview.json").read_text(encoding="utf-8"))
        self.assertIn("table_cell_images_not_walked", {finding["code"] for finding in report["findings"]})


if __name__ == "__main__":
    unittest.main()
