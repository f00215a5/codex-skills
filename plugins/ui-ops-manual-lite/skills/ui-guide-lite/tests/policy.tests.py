from __future__ import annotations

import copy
import hashlib
import json
import subprocess
import sys
import tempfile
import unittest
import zipfile
from pathlib import Path

from PIL import Image
from docx import Document


ROOT = Path(__file__).parents[1]
SCRIPTS = ROOT / "scripts"
ANNOTATE = SCRIPTS / "annotate.py"
REDACT = SCRIPTS / "redact.py"
BUILD = SCRIPTS / "build_docx.py"
VERIFY = SCRIPTS / "verify_docx.py"
AUDIT = SCRIPTS / "audit_delivery.py"


def run(script: Path, *args: str) -> subprocess.CompletedProcess[str]:
    return subprocess.run(
        [sys.executable, str(script), *args],
        check=False,
        capture_output=True,
        text=True,
        encoding="utf-8",
    )


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def minimal_manifest(image_ref: str = "annotated/step.png") -> dict:
    return {
        "title": "訂單管理操作說明書",
        "subtitle": "系統後台操作指引",
        "applicableScreens": ["後台 > 訂單管理"],
        "revision": {"version": "1.0", "date": "2026-01-15", "summary": "初版"},
        "usageReminders": ["以測試資料操作，截圖中的識別資料已遮蔽。"],
        "commonRules": ["必填欄位以紅色 * 標示。"],
        "fontName": "Microsoft JhengHei",
        "chapter": {
            "title": "訂單管理功能",
            "entry": {"caption": "功能入口：訂單管理。", "image": image_ref},
            "sections": [{
                "title": "查詢訂單",
                "steps": [{"caption": "紅框 1：點選查詢。", "image": image_ref}],
                "fields": [{"name": "訂單編號", "definition": "唯一識別碼。", "required": "選填", "limits": "文字"}],
                "impact": "查詢不修改資料。",
                "verification": "確認結果正確。",
            }],
        },
        "updateLog": [{"version": "1.0", "date": "2026-01-15", "changes": "初版"}],
    }


class RedactionAndAnnotationPolicyTests(unittest.TestCase):
    def test_proposed_annotation_can_draw_preview_but_needs_approval_for_build(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            work = Path(directory)
            raw = work / "raw.png"
            Image.new("RGB", (120, 80), "white").save(raw)
            manifest = {
                "sourceImage": "raw.png",
                "originalImageSize": {"width": 120, "height": 80},
                "annotations": [{
                    "id": "1",
                    "controlName": "儲存",
                    "caption": "紅框 1：儲存按鈕。",
                    "bbox": {"x": 20, "y": 20, "width": 30, "height": 20},
                    "status": "proposed",
                    "provenance": "manual-adjusted",
                }],
            }
            ann = work / "annotations.json"
            ann.write_text(json.dumps(manifest, ensure_ascii=False), encoding="utf-8")
            preview = work / "preview.png"

            check = run(ANNOTATE, "check", "--image", str(raw), "--annotations", str(ann))
            draw = run(ANNOTATE, "draw", "--image", str(raw), "--annotations", str(ann), "--output", str(preview))
            approval = run(ANNOTATE, "check", "--require-approved", "--image", str(raw), "--annotations", str(ann))

            self.assertEqual(check.returncode, 0, check.stderr)
            self.assertEqual(draw.returncode, 0, draw.stderr)
            self.assertTrue(preview.is_file())
            self.assertNotEqual(approval.returncode, 0)
            self.assertIn("approval", approval.stderr.lower())

    def test_redaction_draws_opaque_rectangle_and_records_hash_provenance(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            work = Path(directory)
            raw = work / "raw.png"
            redacted = work / "redacted.png"
            Image.new("RGB", (100, 60), "white").save(raw)
            manifest_path = work / "redaction.json"
            manifest_path.write_text(json.dumps({
                "sourceImage": "raw.png",
                "sourceKind": "captured",
                "redactedImage": "redacted.png",
                "originalImageSize": {"width": 100, "height": 60},
                "redactions": [{
                    "category": "policy-number",
                    "bbox": {"x": 10, "y": 12, "width": 30, "height": 16},
                    "method": "opaque-rectangle",
                    "status": "pending",
                }],
            }, ensure_ascii=False), encoding="utf-8")

            draw = run(REDACT, "draw", "--image", str(raw), "--manifest", str(manifest_path), "--output", str(redacted))
            self.assertEqual(draw.returncode, 0, draw.stderr)
            self.assertTrue(redacted.is_file())
            with Image.open(redacted) as image:
                self.assertEqual(image.getpixel((15, 15)), (31, 31, 31))
                self.assertEqual(image.size, (100, 60))
            self.assertIn(sha256(raw), draw.stdout)
            self.assertIn(sha256(redacted), draw.stdout)

    def test_schematic_source_requires_explicit_approval(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            work = Path(directory)
            image = work / "annotated.png"
            Image.new("RGB", (100, 60), "white").save(image)
            manifest = minimal_manifest("annotated.png")
            manifest["screenshotSource"] = {"kind": "schematic"}
            path = work / "manual.json"
            path.write_text(json.dumps(manifest, ensure_ascii=False), encoding="utf-8")
            output = work / "manual.docx"
            blocked = run(BUILD, "--manifest", str(path), "--output", str(output))
            self.assertNotEqual(blocked.returncode, 0)
            self.assertIn("schematic", blocked.stderr.lower())

    def test_redaction_source_kind_cannot_be_inferred_from_screenshot_source(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            work = Path(directory)
            image = work / "raw.png"
            Image.new("RGB", (20, 20), "white").save(image)
            manifest_path = work / "redaction.json"
            manifest_path.write_text(json.dumps({
                "sourceImage": image.name,
                "screenshotSource": {"kind": "captured"},
                "redactions": [{"category": "address", "bbox": {"x": 1, "y": 1, "width": 2, "height": 2}}],
            }, ensure_ascii=False), encoding="utf-8")
            result = run(REDACT, "check", "--image", str(image), "--manifest", str(manifest_path))
            self.assertNotEqual(result.returncode, 0)
            self.assertIn("sourceKind is required", result.stderr)

    def test_checked_redaction_requires_actual_output_and_matching_hashes(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            work = Path(directory)
            raw = work / "raw.png"
            redacted = work / "redacted.png"
            Image.new("RGB", (100, 60), "white").save(raw)
            manifest = {
                "sourceImage": "raw.png",
                "sourceKind": "captured",
                "redactedImage": "redacted.png",
                "originalImageSize": {"width": 100, "height": 60},
                "redactions": [{
                    "category": "policy-number",
                    "bbox": {"x": 10, "y": 12, "width": 30, "height": 16},
                    "method": "opaque-rectangle",
                    "status": "checked",
                }],
            }
            manifest_path = work / "redaction.json"
            manifest_path.write_text(json.dumps(manifest, ensure_ascii=False), encoding="utf-8")
            self.assertEqual(run(REDACT, "draw", "--image", str(raw), "--manifest", str(manifest_path), "--output", str(redacted)).returncode, 0)
            missing_hash = run(REDACT, "check", "--require-checked", "--image", str(raw), "--redacted", str(redacted), "--manifest", str(manifest_path))
            self.assertNotEqual(missing_hash.returncode, 0)
            self.assertIn("Sha256", missing_hash.stderr)
            manifest["sourceSha256"] = sha256(raw)
            manifest["redactedSha256"] = sha256(redacted)
            manifest_path.write_text(json.dumps(manifest, ensure_ascii=False), encoding="utf-8")
            checked = run(REDACT, "check", "--require-checked", "--image", str(raw), "--redacted", str(redacted), "--manifest", str(manifest_path))
            self.assertEqual(checked.returncode, 0, checked.stderr)

    def test_redaction_outputs_cannot_overwrite_raw_or_manifest(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            work = Path(directory)
            raw = work / "raw.png"
            Image.new("RGB", (20, 20), "white").save(raw)
            manifest_path = work / "redaction.json"
            manifest_path.write_text(json.dumps({
                "sourceImage": "raw.png", "sourceKind": "captured",
                "redactions": [{"category": "address", "bbox": {"x": 1, "y": 1, "width": 2, "height": 2}}],
            }), encoding="utf-8")
            same_raw = run(REDACT, "draw", "--image", str(raw), "--manifest", str(manifest_path), "--output", str(raw))
            same_manifest = run(REDACT, "draw", "--image", str(raw), "--manifest", str(manifest_path), "--output", str(manifest_path))
            self.assertEqual(same_raw.returncode, 2)
            self.assertEqual(same_manifest.returncode, 2)

    def test_amount_is_visible_by_default_and_requires_explicit_mask_policy(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            work = Path(directory)
            raw = work / "raw.png"
            Image.new("RGB", (20, 20), "white").save(raw)
            manifest_path = work / "redaction.json"
            payload = {
                "sourceImage": "raw.png", "sourceKind": "captured",
                "redactions": [{"category": "amount", "bbox": {"x": 1, "y": 1, "width": 2, "height": 2}}],
            }
            manifest_path.write_text(json.dumps(payload, ensure_ascii=False), encoding="utf-8")
            blocked = run(REDACT, "check", "--image", str(raw), "--manifest", str(manifest_path))
            self.assertNotEqual(blocked.returncode, 0)
            self.assertIn("amountPolicy", blocked.stderr)
            payload["amountPolicy"] = "mask"
            manifest_path.write_text(json.dumps(payload, ensure_ascii=False), encoding="utf-8")
            allowed = run(REDACT, "check", "--image", str(raw), "--manifest", str(manifest_path))
            self.assertEqual(allowed.returncode, 0, allowed.stderr)


class TableLayoutPolicyTests(unittest.TestCase):
    def test_builder_writes_explicit_center_and_autofit_table_properties(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            work = Path(directory)
            image = work / "annotated" / "step.png"
            image.parent.mkdir()
            Image.new("RGB", (100, 60), "white").save(image)
            manifest_path = work / "manual.json"
            manifest_path.write_text(json.dumps(minimal_manifest(), ensure_ascii=False), encoding="utf-8")
            output = work / "manual.docx"
            built = run(BUILD, "--manifest", str(manifest_path), "--output", str(output))
            self.assertEqual(built.returncode, 0, built.stderr)
            with zipfile.ZipFile(output) as archive:
                xml = archive.read("word/document.xml").decode("utf-8")
            self.assertGreaterEqual(xml.count('w:jc w:val="center"'), 2)
            self.assertIn('w:tblLayout w:type="autofit"', xml)
            self.assertNotIn('w:jc w:val="left"', xml)

            document = Document(output)
            for table in document.tables:
                self.assertEqual(table.alignment, 1)  # WD_TABLE_ALIGNMENT.CENTER

    def test_verify_rejects_left_aligned_table(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            work = Path(directory)
            image = work / "annotated" / "step.png"
            image.parent.mkdir()
            Image.new("RGB", (100, 60), "white").save(image)
            manifest_path = work / "manual.json"
            manifest_path.write_text(json.dumps(minimal_manifest(), ensure_ascii=False), encoding="utf-8")
            output = work / "manual.docx"
            self.assertEqual(run(BUILD, "--manifest", str(manifest_path), "--output", str(output)).returncode, 0)
            document = Document(output)
            for table in document.tables:
                table.alignment = 0  # WD_TABLE_ALIGNMENT.LEFT
            document.save(output)
            verify = run(VERIFY, "--docx", str(output), "--manifest", str(manifest_path))
            self.assertNotEqual(verify.returncode, 0)
            self.assertIn("explicit center alignment", verify.stdout)

    def test_audit_allows_autofit_percentage_width_without_mixing_units(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            work = Path(directory)
            image = work / "annotated" / "step.png"
            image.parent.mkdir()
            Image.new("RGB", (100, 60), "white").save(image)
            manifest_path = work / "manual.json"
            manifest_path.write_text(json.dumps(minimal_manifest(), ensure_ascii=False), encoding="utf-8")
            output = work / "manual.docx"
            self.assertEqual(run(BUILD, "--manifest", str(manifest_path), "--output", str(output)).returncode, 0)
            pct = work / "pct.docx"
            with zipfile.ZipFile(output, "r") as source, zipfile.ZipFile(pct, "w") as destination:
                for item in source.infolist():
                    payload = source.read(item.filename)
                    if item.filename == "word/document.xml":
                        payload = payload.replace(
                            b'<w:tblW w:type="auto" w:w="0"/>',
                            b'<w:tblW w:type="pct" w:w="5000"/>',
                        )
                    destination.writestr(item, payload)
            report_path = work / "audit.json"
            self.assertEqual(run(AUDIT, "--docx", str(pct), "--output", str(report_path)).returncode, 0)
            report = json.loads(report_path.read_text(encoding="utf-8"))
            self.assertEqual(report["tables"][0]["width"]["status"], "pass")
            self.assertEqual(report["tables"][0]["grid"]["status"], "pass")

    def test_audit_rejects_invalid_percentage_width(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            work = Path(directory)
            image = work / "annotated" / "step.png"
            image.parent.mkdir()
            Image.new("RGB", (100, 60), "white").save(image)
            manifest_path = work / "manual.json"
            manifest_path.write_text(json.dumps(minimal_manifest(), ensure_ascii=False), encoding="utf-8")
            output = work / "manual.docx"
            self.assertEqual(run(BUILD, "--manifest", str(manifest_path), "--output", str(output)).returncode, 0)
            invalid = work / "invalid-pct.docx"
            with zipfile.ZipFile(output, "r") as source, zipfile.ZipFile(invalid, "w") as destination:
                for item in source.infolist():
                    payload = source.read(item.filename)
                    if item.filename == "word/document.xml":
                        payload = payload.replace(
                            b'<w:tblW w:type="auto" w:w="0"/>',
                            b'<w:tblW w:type="pct" w:w="6000"/>',
                        )
                    destination.writestr(item, payload)
            report_path = work / "invalid-audit.json"
            self.assertEqual(run(AUDIT, "--docx", str(invalid), "--output", str(report_path)).returncode, 0)
            report = json.loads(report_path.read_text(encoding="utf-8"))
            self.assertEqual(report["tables"][0]["width"]["status"], "fail")

    def test_builder_rejects_raw_image_reference(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            work = Path(directory)
            image = work / "raw" / "screen.png"
            image.parent.mkdir()
            Image.new("RGB", (100, 60), "white").save(image)
            manifest_path = work / "manual.json"
            manifest_path.write_text(json.dumps(minimal_manifest("raw/screen.png"), ensure_ascii=False), encoding="utf-8")
            output = work / "manual.docx"
            build = run(BUILD, "--manifest", str(manifest_path), "--output", str(output))
            self.assertNotEqual(build.returncode, 0)
            self.assertIn("raw or unredacted", build.stderr.lower())


class LiteAuditPolicyTests(unittest.TestCase):
    def _audit_fixture(self, work: Path) -> dict:
        raw_image = work / "raw" / "step.png"
        redacted_image = work / "redacted" / "step.png"
        image = work / "annotated" / "step.png"
        raw_image.parent.mkdir()
        redacted_image.parent.mkdir()
        image.parent.mkdir()
        Image.new("RGB", (100, 60), "white").save(raw_image)
        screenshot_path = work / "screenshot.json"
        screenshot = {
            "schema_version": 1,
            "sourceImage": "raw/step.png",
            "sourceSha256": sha256(raw_image),
            "redactedImage": "redacted/step.png",
            "annotatedImage": "annotated/step.png",
            "sourceKind": "captured",
            "captureKind": "full-page",
            "captureState": {
                "id": "audit-fixture-01",
                "rawSha256": sha256(raw_image),
                "viewportCssSize": {"width": 100, "height": 60},
                "scroll": {"x": 0, "y": 0},
                "calibration": {
                    "mode": "known-control",
                    "pngSize": {"width": 100, "height": 60},
                    "viewportCssSize": {"width": 100, "height": 60},
                    "knownControl": {"bbox": {"x": 20, "y": 15, "width": 20, "height": 15}},
                },
            },
            "originalImageSize": {"width": 100, "height": 60},
            "protectedAreas": [],
            "redactions": [{
                "category": "policy-number",
                "bbox": {"x": 1, "y": 1, "width": 2, "height": 2},
                "method": "opaque-rectangle",
                "status": "checked",
            }],
            "annotations": [{
                "id": "1",
                "controlName": "查詢",
                "caption": "紅框 1：查詢按鈕。",
                "bbox": {"x": 20, "y": 15, "width": 20, "height": 15},
                "source": "dom-derived",
                "provenance": "dom-derived",
                "status": "verified",
            }],
            "reviewStatus": "checked",
        }
        screenshot_path.write_text(json.dumps(screenshot, ensure_ascii=False), encoding="utf-8")
        redaction = run(
            REDACT,
            "draw",
            "--image",
            str(raw_image),
            "--manifest",
            str(screenshot_path),
            "--output",
            str(redacted_image),
            "--provenance-output",
            str(work / "redaction-provenance.json"),
        )
        self.assertEqual(redaction.returncode, 0, redaction.stderr)
        screenshot["redactedSha256"] = sha256(redacted_image)
        screenshot["redactionProvenance"] = "redaction-provenance.json"
        screenshot_path.write_text(json.dumps(screenshot, ensure_ascii=False), encoding="utf-8")
        annotation = run(
            ANNOTATE,
            "draw",
            "--image",
            str(redacted_image),
            "--annotations",
            str(screenshot_path),
            "--output",
            str(image),
            "--provenance-output",
            str(work / "annotation-provenance.json"),
        )
        self.assertEqual(annotation.returncode, 0, annotation.stderr)
        screenshot["annotatedSha256"] = sha256(image)
        screenshot["annotationProvenance"] = "annotation-provenance.json"
        screenshot_path.write_text(json.dumps(screenshot, ensure_ascii=False), encoding="utf-8")
        manifest_path = work / "手冊-manual.json"
        manifest_path.write_text(json.dumps(minimal_manifest(), ensure_ascii=False), encoding="utf-8")
        output = work / "manual.docx"
        built = run(BUILD, "--manifest", str(manifest_path), "--output", str(output))
        self.assertEqual(built.returncode, 0, built.stderr)
        requirements = work / "需求範圍.md"
        requirements.write_text("Synthetic test requirements only.", encoding="utf-8")
        build_manifest = work / "建置-manual.json"
        build_manifest.write_bytes(manifest_path.read_bytes())
        image_name = image.relative_to(work).as_posix()
        requirements_name = requirements.name
        manifest_name = build_manifest.name
        return {
            "work": work,
            "docx": output,
            "image": image,
            "foreign_image": work / "foreign.png",
            "requirements": requirements,
            "build_manifest": build_manifest,
            "review": {
                "schema_version": 1,
                "artifact": {"sha256": sha256(output)},
                "builder": {"id": "builder:test"},
                "reviewer": {"id": "reviewer:other"},
                "reviewed_files": [
                    {"kind": "image", "path": image_name, "sha256": sha256(image)},
                    {"kind": "support", "role": "build_manifest", "path": manifest_name, "sha256": sha256(build_manifest)},
                    {"kind": "support", "role": "requirements", "path": requirements_name, "sha256": sha256(requirements)},
                ],
                "checks": {
                    key: {
                        "status": "pass",
                        "reason": "Synthetic schema test, no actual semantic approval.",
                        "evidence": [image_name] if key in {"redaction", "operation"}
                        else [requirements_name, manifest_name],
                    }
                    for key in ("requirements", "redaction", "operation", "layout_structure")
                },
                "render_visual": {"status": "out_of_scope", "reason": "Lite profile has no renderer."},
                "overall": "pass",
            },
        }

    def _run_review_case(self, fixture: dict, name: str, payload: dict, *, manifest: Path | None = None) -> dict:
        work = fixture["work"]
        review_path = work / f"review-{name}.json"
        report_path = work / f"audit-{name}.json"
        review_path.write_text(json.dumps(payload, ensure_ascii=False), encoding="utf-8")
        args = ["--docx", str(fixture["docx"]), "--review", str(review_path), "--output", str(report_path)]
        if manifest is not None:
            args.extend(["--manifest", str(manifest)])
        result = run(AUDIT, *args)
        self.assertEqual(result.returncode, 0, result.stdout + result.stderr)
        return json.loads(report_path.read_text(encoding="utf-8"))["independent_review"]

    def test_lite_audit_rejects_unbound_or_stale_review_records(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            fixture = self._audit_fixture(Path(directory))
            valid = self._run_review_case(fixture, "valid-localized-role-files", fixture["review"])
            self.assertEqual(valid["status"], "pass")

            def stale_image(payload: dict) -> None:
                payload["reviewed_files"][0]["sha256"] = "0" * 64

            def foreign_image_hash(payload: dict) -> None:
                Image.new("RGB", (100, 60), "#eeeeee").save(fixture["foreign_image"])
                payload["reviewed_files"][0]["path"] = fixture["foreign_image"].name
                payload["reviewed_files"][0]["sha256"] = sha256(fixture["foreign_image"])
                for key in ("redaction", "operation"):
                    payload["checks"][key]["evidence"] = [fixture["foreign_image"].name]

            def stale_manifest(payload: dict) -> None:
                payload["reviewed_files"][1]["sha256"] = "0" * 64

            def unrelated_redaction_evidence(payload: dict) -> None:
                payload["checks"]["redaction"]["evidence"] = [fixture["requirements"].name]

            cases: list[tuple[str, object]] = [
                ("self-review", lambda p: p["reviewer"].update(id="builder:test")),
                ("pending-reviewer", lambda p: p["reviewer"].update(id="pending")),
                ("stale-docx", lambda p: p["artifact"].update(sha256="0" * 64)),
                ("stale-image", stale_image),
                ("foreign-image-hash", foreign_image_hash),
                ("stale-build-manifest", stale_manifest),
                ("unrelated-redaction-evidence", unrelated_redaction_evidence),
                ("missing-requirements-role", lambda p: p["reviewed_files"][2].pop("role")),
                ("duplicate-reviewed-file", lambda p: p["reviewed_files"].append(copy.deepcopy(p["reviewed_files"][0]))),
                ("false-na", lambda p: p["checks"]["redaction"].update(status="na")),
                ("invalid-schema-boolean", lambda p: p.update(schema_version=True)),
            ]
            for name, mutate in cases:
                with self.subTest(case=name):
                    changed = copy.deepcopy(fixture["review"])
                    mutate(changed)
                    report = self._run_review_case(fixture, name, changed)
                    self.assertNotEqual(report["status"], "pass")

            missing_review_report = fixture["work"] / "audit-missing-review.json"
            missing = run(AUDIT, "--docx", str(fixture["docx"]), "--output", str(missing_review_report))
            self.assertEqual(missing.returncode, 0, missing.stdout + missing.stderr)
            missing_report = json.loads(missing_review_report.read_text(encoding="utf-8"))
            self.assertEqual(missing_report["independent_review"]["status"], "blocked")
            self.assertIn("review_missing", missing_report["independent_review"]["reason_code"])

            changed_manifest = fixture["work"] / "different-manual.json"
            changed_manifest.write_text(json.dumps({"changed": True}), encoding="utf-8")
            explicit_mismatch = self._run_review_case(
                fixture, "different-explicit-manifest", fixture["review"], manifest=changed_manifest
            )
            self.assertNotEqual(explicit_mismatch["status"], "pass")

    def test_lite_review_allows_render_out_of_scope_but_requires_hash_bound_evidence(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            work = Path(directory)
            image = work / "annotated" / "step.png"
            image.parent.mkdir()
            Image.new("RGB", (100, 60), "white").save(image)
            manifest_path = work / "manual.json"
            manifest_path.write_text(json.dumps(minimal_manifest(), ensure_ascii=False), encoding="utf-8")
            output = work / "manual.docx"
            self.assertEqual(run(BUILD, "--manifest", str(manifest_path), "--output", str(output)).returncode, 0)
            requirements = work / "requirements.md"
            requirements.write_text("已確認需求。", encoding="utf-8")
            build_manifest = work / "build-manifest.json"
            build_manifest.write_bytes(manifest_path.read_bytes())
            review = work / "review.json"
            reviewed = [
                {"kind": "image", "path": "annotated/step.png", "sha256": sha256(image)},
                {"kind": "support", "role": "build-manifest", "path": "build-manifest.json", "sha256": sha256(build_manifest)},
                {"kind": "support", "role": "requirements", "path": "requirements.md", "sha256": sha256(requirements)},
            ]
            checks = {
                key: {
                    "status": "pass",
                    "reason": "已核對。",
                    "evidence": ["requirements.md"] if key in {"requirements", "layout_structure"}
                    else ["annotated/step.png"],
                }
                for key in ("requirements", "redaction", "operation", "layout_structure")
            }
            payload = {
                "schema_version": 1,
                "artifact": {"sha256": sha256(output)},
                "builder": {"id": "builder:test"},
                "reviewer": {"id": "reviewer:other"},
                "reviewed_files": reviewed,
                "checks": checks,
                "render_visual": {"status": "out_of_scope", "reason": "輕量版不提供 DOCX renderer。"},
                "overall": "pass",
            }
            review.write_text(json.dumps(payload, ensure_ascii=False), encoding="utf-8")
            report_path = work / "audit.json"
            result = run(AUDIT, "--docx", str(output), "--review", str(review), "--output", str(report_path))
            self.assertEqual(result.returncode, 0, result.stdout + result.stderr)
            report = json.loads(report_path.read_text(encoding="utf-8"))
            self.assertEqual(report["independent_review"]["status"], "pass")
            self.assertEqual(report["independent_review"]["render_visual"]["status"], "out_of_scope")
            self.assertNotIn("overall", report)

    def test_lite_audit_rejects_render_visual_pass(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            work = Path(directory)
            image = work / "annotated" / "step.png"
            image.parent.mkdir()
            Image.new("RGB", (100, 60), "white").save(image)
            manifest_path = work / "manual.json"
            manifest_path.write_text(json.dumps(minimal_manifest(), ensure_ascii=False), encoding="utf-8")
            output = work / "manual.docx"
            self.assertEqual(run(BUILD, "--manifest", str(manifest_path), "--output", str(output)).returncode, 0)
            requirements = work / "requirements.md"
            requirements.write_text("已確認需求。", encoding="utf-8")
            build_manifest = work / "build-manifest.json"
            build_manifest.write_bytes(manifest_path.read_bytes())
            review = work / "review.json"
            reviewed = [
                {"kind": "image", "path": "annotated/step.png", "sha256": sha256(image)},
                {"kind": "support", "role": "build-manifest", "path": "build-manifest.json", "sha256": sha256(build_manifest)},
                {"kind": "support", "role": "requirements", "path": "requirements.md", "sha256": sha256(requirements)},
            ]
            checks = {
                key: {
                    "status": "pass",
                    "reason": "已核對。",
                    "evidence": ["requirements.md"] if key in {"requirements", "layout_structure"}
                    else ["annotated/step.png"],
                }
                for key in ("requirements", "redaction", "operation", "layout_structure")
            }
            review.write_text(json.dumps({
                "schema_version": 1,
                "artifact": {"sha256": sha256(output)},
                "builder": {"id": "builder:test"},
                "reviewer": {"id": "reviewer:other"},
                "reviewed_files": reviewed,
                "checks": checks,
                "render_visual": {"status": "pass", "reason": "錯誤宣稱。", "evidence": ["requirements.md"]},
                "overall": "pass",
            }, ensure_ascii=False), encoding="utf-8")
            report_path = work / "audit.json"
            result = run(AUDIT, "--docx", str(output), "--review", str(review), "--output", str(report_path))
            self.assertEqual(result.returncode, 0)
            report = json.loads(report_path.read_text(encoding="utf-8"))
            self.assertNotEqual(report["independent_review"]["status"], "pass")
            self.assertIn("render_visual_must_not_pass", report["independent_review"]["reason_codes"])

    def test_lite_audit_binds_each_media_image_to_redaction_annotation_provenance(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            fixture = self._audit_fixture(Path(directory))
            build_payload = json.loads(fixture["build_manifest"].read_text(encoding="utf-8"))
            build_payload["chapter"]["entry"]["evidence"] = {
                "redactionManifest": "screenshot.json",
                "annotationManifest": "screenshot.json",
                "provenance": "redaction-provenance.json",
                "annotationProvenance": "annotation-provenance.json",
            }
            fixture["build_manifest"].write_text(json.dumps(build_payload), encoding="utf-8")
            payload = copy.deepcopy(fixture["review"])
            payload["reviewed_files"].append({
                "kind": "support", "role": "screenshot_manifest", "path": "screenshot.json",
                "sha256": sha256(fixture["work"] / "screenshot.json"),
            })
            payload["reviewed_files"].append({
                "kind": "support", "role": "redaction_provenance", "path": "redaction-provenance.json",
                "sha256": sha256(fixture["work"] / "redaction-provenance.json"),
            })
            payload["reviewed_files"].append({
                "kind": "support", "role": "annotation_provenance", "path": "annotation-provenance.json",
                "sha256": sha256(fixture["work"] / "annotation-provenance.json"),
            })
            for item in payload["reviewed_files"]:
                if item.get("path") == fixture["build_manifest"].name:
                    item["sha256"] = sha256(fixture["build_manifest"])
            payload["imageEvidence"] = [{
                "image": "annotated/step.png",
                "redactionManifest": "screenshot.json",
                "annotationManifest": "screenshot.json",
                "provenance": "redaction-provenance.json",
                "annotationProvenance": "annotation-provenance.json",
            }]
            valid = self._run_review_case(
                fixture, "per-image-chain", payload, manifest=fixture["build_manifest"]
            )
            self.assertEqual(valid["status"], "pass")
            self.assertEqual(valid["image_evidence_chain"]["status"], "pass")

            broken = copy.deepcopy(payload)
            broken["imageEvidence"][0].pop("provenance")
            blocked = self._run_review_case(
                fixture, "per-image-chain-missing-provenance", broken, manifest=fixture["build_manifest"]
            )
            self.assertNotEqual(blocked["status"], "pass")
            self.assertIn("image_evidence_chain_provenance_missing", blocked["reason_codes"])

            broken_annotation = copy.deepcopy(payload)
            broken_annotation["imageEvidence"][0].pop("annotationProvenance")
            blocked_annotation = self._run_review_case(
                fixture, "per-image-chain-missing-annotation-provenance", broken_annotation,
                manifest=fixture["build_manifest"],
            )
            self.assertNotEqual(blocked_annotation["status"], "pass")

            unrelated = copy.deepcopy(payload)
            unrelated["imageEvidence"][0]["provenance"] = "需求範圍.md"
            unrelated["imageEvidence"][0]["annotationProvenance"] = "需求範圍.md"
            blocked_unrelated = self._run_review_case(
                fixture, "per-image-chain-unrelated-support", unrelated,
                manifest=fixture["build_manifest"],
            )
            self.assertNotEqual(blocked_unrelated["status"], "pass")
            self.assertTrue(any("not_manifest_provenance" in code for code in blocked_unrelated["reason_codes"]))


if __name__ == "__main__":
    unittest.main()
