from __future__ import annotations

import json
import subprocess
import sys
import tempfile
import unittest
import zipfile
import xml.etree.ElementTree as ET
from pathlib import Path

from PIL import Image
from docx import Document
from docx.oxml.ns import qn

SCRIPTS = Path(__file__).parents[1] / "scripts"
BUILD_SCRIPT = SCRIPTS / "build_docx.py"
VERIFY_SCRIPT = SCRIPTS / "verify_docx.py"

MINIMAL_MANIFEST = {
    "title": "訂單管理操作說明書",
    "subtitle": "系統後台操作指引",
    "applicableScreens": ["系統後台 > 訂單管理 > 訂單列表"],
    "revision": {"version": "1.0", "date": "2026-01-15", "summary": "初版"},
    "usageReminders": ["截圖與操作以測試環境為準，遮蔽敏感資料。"],
    "commonRules": ["必填欄位以紅色 * 標示。"],
    "fontName": "PingFang TC",
    "chapter": {
        "title": "訂單管理功能",
        "entry": {"caption": "功能入口：訂單管理。", "image": "annotated/entry.png",
                  "detail": "從側邊欄進入。"},
        "sections": [
            {
                "title": "查詢訂單",
                "preconditions": "已登入系統後台。",
                "steps": [
                    {"caption": "紅框 1：點選「訂單查詢」。", "image": "annotated/step1.png",
                     "detail": "系統顯示查詢表單。"},
                    {"caption": "紅框 2：輸入訂單編號後按查詢。", "image": "annotated/step2.png"},
                ],
                "fields": [
                    {"name": "訂單編號", "definition": "訂單的唯一識別碼。",
                     "required": "選填", "limits": "最大 30 字元", "display": "查詢區"},
                ],
                "impact": "查詢結果會依權限過濾，不影響資料本身。",
                "failure": "無符合資料時顯示空清單與提示。",
                "verification": "以同一筆訂單編號重查可得到相同結果。",
            },
            {
                "title": "開立訂單",
                "steps": [
                    {"caption": "紅框 1：點選「新增訂單」。", "image": "annotated/step3.png"},
                ],
                "fields": [
                    {"name": "客戶", "definition": "訂單所屬客戶。",
                     "required": "必填", "limits": "下拉選單"},
                ],
                "impact": "成功後建立新訂單並產生訂單編號。",
                "verification": "於訂單列表可見新訂單。",
            },
        ],
    },
    "updateLog": [{"version": "1.0", "date": "2026-01-15", "changes": "初版"}],
}


def make_annotated_image(path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    Image.new("RGB", (320, 200), "#F7FFFC").save(path)


def run(script: Path, *arguments: str) -> subprocess.CompletedProcess[str]:
    return subprocess.run(
        [sys.executable, str(script), *arguments],
        check=False,
        capture_output=True,
        text=True,
        encoding="utf-8",
    )


class BuildAndVerifyTests(unittest.TestCase):
    def setUp(self) -> None:
        self.temp_dir = tempfile.TemporaryDirectory()
        self.workspace = Path(self.temp_dir.name)
        self.images = self.workspace / "annotated"
        for name in ("entry.png", "step1.png", "step2.png", "step3.png"):
            make_annotated_image(self.images / name)
        self.manifest_path = self.workspace / "manual.json"
        self.manifest_path.write_text(json.dumps(MINIMAL_MANIFEST), encoding="utf-8")
        self.docx_path = self.workspace / "manual.docx"

    def tearDown(self) -> None:
        self.temp_dir.cleanup()

    def build_and_verify(self) -> subprocess.CompletedProcess[str]:
        build = run(BUILD_SCRIPT, "--manifest", str(self.manifest_path), "--output", str(self.docx_path))
        self.assertEqual(build.returncode, 0, build.stderr)
        return run(VERIFY_SCRIPT, "--docx", str(self.docx_path), "--manifest", str(self.manifest_path))

    def test_round_trip_build_passes_structure_qa(self) -> None:
        verify = self.build_and_verify()
        self.assertEqual(verify.returncode, 0, verify.stdout + verify.stderr)
        self.assertIn("all structure checks passed", verify.stdout)

    def test_missing_image_fails_build(self) -> None:
        manifest = json.loads(self.manifest_path.read_text(encoding="utf-8"))
        manifest["chapter"]["sections"][0]["steps"][0]["image"] = "annotated/does-not-exist.png"
        self.manifest_path.write_text(json.dumps(manifest), encoding="utf-8")

        build = run(BUILD_SCRIPT, "--manifest", str(self.manifest_path), "--output", str(self.docx_path))

        self.assertNotEqual(build.returncode, 0)
        self.assertIn("image not found", build.stderr)

    def test_runtime_availability_warning_is_rejected(self) -> None:
        manifest = json.loads(self.manifest_path.read_text(encoding="utf-8"))
        manifest["usageReminders"] = ["word-render 未安裝，LibreOffice 不可用。"]
        self.manifest_path.write_text(json.dumps(manifest), encoding="utf-8")

        build = run(BUILD_SCRIPT, "--manifest", str(self.manifest_path), "--output", str(self.docx_path))

        self.assertNotEqual(build.returncode, 0)
        self.assertIn("runtime availability warning", build.stderr)
        self.assertFalse(self.docx_path.exists())

    def test_runtime_warning_guard_allows_password_content(self) -> None:
        manifest = json.loads(self.manifest_path.read_text(encoding="utf-8"))
        manifest["usageReminders"] = ["Password unavailable 時請依系統錯誤訊息重設密碼。"]
        self.manifest_path.write_text(json.dumps(manifest), encoding="utf-8")

        build = run(BUILD_SCRIPT, "--manifest", str(self.manifest_path), "--output", str(self.docx_path))

        self.assertEqual(build.returncode, 0, build.stderr)
        self.assertTrue(self.docx_path.is_file())

    def test_verify_catches_empty_update_log(self) -> None:
        manifest = json.loads(self.manifest_path.read_text(encoding="utf-8"))
        manifest["updateLog"] = []
        self.manifest_path.write_text(json.dumps(manifest), encoding="utf-8")
        build = run(BUILD_SCRIPT, "--manifest", str(self.manifest_path), "--output", str(self.docx_path))
        self.assertEqual(build.returncode, 0, build.stderr)

        verify = run(VERIFY_SCRIPT, "--docx", str(self.docx_path), "--manifest", str(self.manifest_path))

        self.assertNotEqual(verify.returncode, 0)
        self.assertIn("update log has data rows", verify.stdout)

    def test_docx_is_a_valid_zip_package(self) -> None:
        self.build_and_verify()
        import zipfile
        with zipfile.ZipFile(self.docx_path) as archive:
            bad = archive.testzip()
        self.assertIsNone(bad)

    def test_update_log_is_the_first_heading_and_table_after_the_title(self) -> None:
        build = run(BUILD_SCRIPT, "--manifest", str(self.manifest_path), "--output", str(self.docx_path))
        self.assertEqual(build.returncode, 0, build.stderr)

        document = Document(self.docx_path)
        headings = [
            paragraph.text
            for paragraph in document.paragraphs
            if paragraph.style.name.startswith("Heading")
        ]
        self.assertEqual(headings[:2], ["更新紀錄", "修訂狀態"])
        self.assertEqual(
            [cell.text for cell in document.tables[0].rows[0].cells][:3],
            ["版本", "日期", "更新內容"],
        )

    def test_chapters_list_builds_and_verifies_each_chapter(self) -> None:
        manifest = json.loads(self.manifest_path.read_text(encoding="utf-8"))
        chapter = manifest.pop("chapter")
        second = json.loads(json.dumps(chapter))
        second["title"] = "付款管理功能"
        second["entry"]["caption"] = "功能入口：付款管理。"
        manifest["chapters"] = [chapter, second]
        self.manifest_path.write_text(json.dumps(manifest, ensure_ascii=False), encoding="utf-8")

        verify = self.build_and_verify()

        self.assertEqual(verify.returncode, 0, verify.stdout + verify.stderr)
        document = Document(self.docx_path)
        headings = [paragraph.text for paragraph in document.paragraphs if paragraph.style.name.startswith("Heading")]
        self.assertIn("訂單管理功能", headings)
        self.assertIn("付款管理功能", headings)

    def test_numbering_definitions_are_unique_and_reset_with_override(self) -> None:
        self.build_and_verify()
        with zipfile.ZipFile(self.docx_path) as archive:
            numbering = archive.read("word/numbering.xml").decode("utf-8")
        self.assertGreaterEqual(numbering.count("<w:abstractNum "), 3)
        self.assertGreaterEqual(numbering.count("<w:num "), 3)
        self.assertIn("<w:startOverride w:val=\"1\"/>", numbering)
        self.assertIn("<w:suff w:val=\"space\"/><w:lvlText", numbering)

    def test_image_and_table_flow_properties_are_explicit(self) -> None:
        self.build_and_verify()
        document = Document(self.docx_path)
        image_paragraphs = [paragraph for paragraph in document.paragraphs if paragraph._p.xpath(".//a:blip")]
        self.assertTrue(image_paragraphs)
        for paragraph in image_paragraphs:
            p_pr = paragraph._p.find(qn("w:pPr"))
            self.assertIsNotNone(p_pr)
            self.assertIsNotNone(p_pr.find(qn("w:keepNext")))
        self.assertTrue(document.tables)
        for table in document.tables:
            header_pr = table.rows[0]._tr.find(qn("w:trPr"))
            self.assertIsNotNone(header_pr)
            self.assertIsNotNone(header_pr.find(qn("w:tblHeader")))
            for row in table.rows:
                row_pr = row._tr.find(qn("w:trPr"))
                self.assertIsNotNone(row_pr)
                self.assertIsNotNone(row_pr.find(qn("w:cantSplit")))

    def test_verify_rejects_final_inline_extent_over_container(self) -> None:
        self.build_and_verify()
        tampered = self.workspace / "oversized-inline.docx"
        namespace = {
            "w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main",
            "wp": "http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing",
        }
        with zipfile.ZipFile(self.docx_path, "r") as source, zipfile.ZipFile(tampered, "w") as destination:
            for item in source.infolist():
                payload = source.read(item.filename)
                if item.filename == "word/document.xml":
                    root = ET.fromstring(payload)
                    extent = root.find(".//wp:inline/wp:extent", namespace)
                    self.assertIsNotNone(extent)
                    extent.set("cx", "999999999")
                    ET.register_namespace("w", namespace["w"])
                    ET.register_namespace("wp", namespace["wp"])
                    payload = ET.tostring(root, encoding="utf-8", xml_declaration=True)
                destination.writestr(item, payload)
        verify = run(VERIFY_SCRIPT, "--docx", str(tampered), "--manifest", str(self.manifest_path))
        self.assertNotEqual(verify.returncode, 0)
        self.assertIn("inline image extents", verify.stdout)

    def test_verify_reports_exact_row_height_as_layout_risk(self) -> None:
        self.build_and_verify()
        tampered = self.workspace / "exact-row-height.docx"
        namespace = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
        with zipfile.ZipFile(self.docx_path, "r") as source, zipfile.ZipFile(tampered, "w") as destination:
            for item in source.infolist():
                payload = source.read(item.filename)
                if item.filename == "word/document.xml":
                    root = ET.fromstring(payload)
                    row = root.find(".//w:tbl/w:tr", namespace)
                    self.assertIsNotNone(row)
                    tr_pr = row.find("w:trPr", namespace)
                    if tr_pr is None:
                        tr_pr = ET.Element(qn("w:trPr"))
                        row.insert(0, tr_pr)
                    height = ET.SubElement(tr_pr, qn("w:trHeight"))
                    height.set(qn("w:val"), "100")
                    height.set(qn("w:hRule"), "exact")
                    ET.register_namespace("w", namespace["w"])
                    payload = ET.tostring(root, encoding="utf-8", xml_declaration=True)
                destination.writestr(item, payload)
        verify = run(VERIFY_SCRIPT, "--docx", str(tampered), "--manifest", str(self.manifest_path))
        self.assertEqual(verify.returncode, 0, verify.stdout or verify.stderr or "")
        self.assertIn("[RISK] table rows use exact fixed heights", verify.stdout)

    def test_verify_rejects_operation_list_without_start_override(self) -> None:
        self.build_and_verify()
        document = Document(self.docx_path)
        operation_num_ids: list[str] = []
        for paragraph in document.paragraphs:
            num_pr = paragraph._p.find(qn("w:pPr") + "/" + qn("w:numPr"))
            if num_pr is None:
                continue
            num_id = num_pr.find(qn("w:numId"))
            if num_id is not None and num_id.get(qn("w:val")) and "紅框" in paragraph.text:
                operation_num_ids.append(num_id.get(qn("w:val")))
        self.assertTrue(operation_num_ids)
        target_num_id = operation_num_ids[0]
        tampered = self.workspace / "tampered.docx"
        with zipfile.ZipFile(self.docx_path, "r") as source, zipfile.ZipFile(tampered, "w") as destination:
            for item in source.infolist():
                payload = source.read(item.filename)
                if item.filename == "word/numbering.xml":
                    root = ET.fromstring(payload)
                    namespace = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
                    target = next(
                        node for node in root.findall("w:num", namespace)
                        if node.get(qn("w:numId")) == target_num_id
                    )
                    override = target.find("w:lvlOverride", namespace)
                    self.assertIsNotNone(override)
                    start = override.find("w:startOverride", namespace)
                    self.assertIsNotNone(start)
                    override.remove(start)
                    ET.register_namespace("w", namespace["w"])
                    payload = ET.tostring(root, encoding="utf-8", xml_declaration=True)
                destination.writestr(item, payload)
        verify = run(VERIFY_SCRIPT, "--docx", str(tampered), "--manifest", str(self.manifest_path))
        self.assertNotEqual(verify.returncode, 0)
        self.assertIn("startOverride", verify.stdout)

    def test_tall_image_fails_with_continuation_guidance(self) -> None:
        tall = self.images / "step1.png"
        Image.new("RGB", (300, 1800), "#F7FFFC").save(tall)

        build = run(BUILD_SCRIPT, "--manifest", str(self.manifest_path), "--output", str(self.docx_path))

        self.assertNotEqual(build.returncode, 0)
        self.assertIn("continuation", build.stderr.lower())


if __name__ == "__main__":
    unittest.main()
