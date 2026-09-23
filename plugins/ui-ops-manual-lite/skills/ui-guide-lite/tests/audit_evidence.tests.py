from __future__ import annotations

import hashlib
import json
import sys
import tempfile
import unittest
from pathlib import Path

from docx import Document
from PIL import Image


SCRIPT_ROOT = Path(__file__).resolve().parents[1] / "scripts"
sys.path.insert(0, str(SCRIPT_ROOT))

import audit_delivery  # noqa: E402


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


class AuditEvidenceTests(unittest.TestCase):
    def _fixture(self, root: Path) -> tuple[Path, Path, Path]:
        delivered = root / "delivered"
        delivered.mkdir()
        image = delivered / "step.png"
        Image.new("RGB", (120, 80), "white").save(image, format="PNG")

        docx = root / "manual.docx"
        document = Document()
        table = document.add_table(rows=1, cols=3)
        for cell, value in zip(table.rows[0].cells, ("版本", "日期", "更新內容")):
            cell.text = value
        document.add_picture(str(image))
        document.save(docx)

        evidence = root / "delivered-images.json"
        evidence.write_text(
            json.dumps({"schema_version": 1, "assets": [{"path": "delivered/step.png", "sha256": sha256(image)}]}),
            encoding="utf-8",
        )
        return docx, image, evidence

    def test_image_evidence_binds_current_media_and_default_layout(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            docx, _image, evidence = self._fixture(Path(directory))
            report = audit_delivery.audit_docx(
                docx,
                require_default_layout=True,
                image_evidence_path=evidence,
            )
            self.assertEqual(report["image_evidence"]["status"], "pass")
            self.assertEqual(report["default_layout"]["status"], "pass")
            self.assertIn(report["mechanical_status"], {"pass", "manual_review"})

    def test_image_evidence_blocks_old_embedded_media_and_protects_asset(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            docx, image, evidence = self._fixture(Path(directory))
            frozen_hash = sha256(image)
            image.write_bytes(b"changed-after-freeze")
            evidence.write_text(
                json.dumps({"schema_version": 1, "assets": [{"path": "delivered/step.png", "sha256": frozen_hash}]}),
                encoding="utf-8",
            )
            report = audit_delivery.audit_docx(docx, image_evidence_path=evidence)
            self.assertEqual(report["image_evidence"]["status"], "fail")
            self.assertIn("image_evidence_asset_hash_mismatch", report["image_evidence"]["reason_codes"])

            # The allowlist now points at a valid new asset, but the DOCX still
            # embeds the old frozen bytes; the media binding must catch that.
            evidence.write_text(
                json.dumps({"schema_version": 1, "assets": [{"path": "delivered/step.png", "sha256": sha256(image)}]}),
                encoding="utf-8",
            )
            stale_docx_report = audit_delivery.audit_docx(docx, image_evidence_path=evidence)
            self.assertEqual(stale_docx_report["image_evidence"]["status"], "fail")
            self.assertIn(
                "image_evidence_embedded_media_not_allowlisted",
                stale_docx_report["image_evidence"]["reason_codes"],
            )

            before = evidence.read_bytes()
            self.assertTrue(
                audit_delivery._output_conflicts(
                    str(image), str(docx), None, None, str(evidence)
                )
            )
            self.assertEqual(evidence.read_bytes(), before)


if __name__ == "__main__":
    unittest.main()
