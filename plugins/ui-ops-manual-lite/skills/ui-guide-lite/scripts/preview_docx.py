#!/usr/bin/env python3
"""Create a content/embedded-image preview from a final lite DOCX.

The helper reads the OOXML package through python-docx, preserves body order
for paragraphs and tables, and scales each embedded image to its inline EMU
dimensions at the requested DPI.  It does not render Word layout, paginate,
or claim font/line-break fidelity.
"""

from __future__ import annotations

import argparse
import hashlib
import html
import io
import json
import math
import sys
import zipfile
from pathlib import Path
from typing import Any, Iterable

from docx import Document
from docx.document import Document as DocumentType
from docx.table import Table, _Cell
from docx.text.paragraph import Paragraph
from docx.oxml.ns import qn
from PIL import Image


EMU_PER_INCH = 914400
SCHEMA_VERSION = 1
MAX_DPI = 600.0
MAX_PREVIEW_SIDE_PX = 20000
MAX_PREVIEW_PIXELS = 50_000_000


def parse_args(argv: list[str] | None = None) -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--docx", required=True, type=Path)
    parser.add_argument("--output-dir", required=True, type=Path)
    parser.add_argument("--dpi", default=96.0, type=float)
    return parser.parse_args(argv)


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _local_name(element: Any) -> str:
    return str(element.tag).rsplit("}", 1)[-1]


def iter_block_items(parent: DocumentType | _Cell) -> Iterable[Paragraph | Table]:
    """Yield paragraphs and tables in the order stored in the body/cell."""

    parent_element = parent.element.body if isinstance(parent, DocumentType) else parent._tc
    parent_obj = parent if isinstance(parent, DocumentType) else parent
    for child in parent_element.iterchildren():
        if _local_name(child) == "p":
            yield Paragraph(child, parent_obj)
        elif _local_name(child) == "tbl":
            yield Table(child, parent_obj)


def _ancestor(element: Any, local_name: str) -> Any | None:
    current = element
    while current is not None:
        if _local_name(current) == local_name:
            return current
        current = current.getparent()
    return None


def _inline_extent(blip: Any) -> tuple[int, int] | None:
    inline = _ancestor(blip, "inline")
    if inline is None:
        return None
    extent = inline.find(qn("wp:extent"))
    if extent is None:
        return None
    try:
        width = int(extent.get("cx"))
        height = int(extent.get("cy"))
    except (TypeError, ValueError):
        return None
    return (width, height) if width > 0 and height > 0 else None


def _image_occurrences(paragraph: Paragraph, document: DocumentType) -> list[dict[str, Any]]:
    occurrences: list[dict[str, Any]] = []
    for blip in paragraph._p.xpath(".//a:blip"):
        relation_id = blip.get(qn("r:embed"))
        if not relation_id or relation_id not in document.part.related_parts:
            occurrences.append({"status": "blocked", "reason": "embedded_relationship_missing"})
            continue
        part = document.part.related_parts[relation_id]
        extent = _inline_extent(blip)
        item: dict[str, Any] = {
            "relationship": relation_id,
            "part": str(getattr(part, "partname", "")),
        }
        if extent is None:
            item.update(status="blocked", reason="inline_extent_missing")
            occurrences.append(item)
            continue
        item["inline_emu"] = {"cx": extent[0], "cy": extent[1]}
        item["status"] = "pass"
        occurrences.append(item)
    return occurrences


def _scale_image(part: Any, destination: Path, extent: tuple[int, int], dpi: float) -> dict[str, Any]:
    if dpi <= 0 or not math.isfinite(dpi) or dpi > MAX_DPI:
        raise ValueError("dpi must be a positive finite number")
    blob = getattr(part, "blob", None)
    if not isinstance(blob, bytes):
        raise ValueError("embedded image has no readable bytes")
    source_blob_sha256 = hashlib.sha256(blob).hexdigest()
    width_px = round(extent[0] / EMU_PER_INCH * dpi)
    height_px = round(extent[1] / EMU_PER_INCH * dpi)
    if (
        width_px <= 0
        or height_px <= 0
        or width_px > MAX_PREVIEW_SIDE_PX
        or height_px > MAX_PREVIEW_SIDE_PX
        or width_px * height_px > MAX_PREVIEW_PIXELS
    ):
        raise ValueError("inline image dimensions exceed the safe preview limit")
    with Image.open(io.BytesIO(blob)) as source:
        source.load()
        source_size = source.size
        image = source.convert("RGBA")
        if image.size != (width_px, height_px):
            image = image.resize((width_px, height_px), Image.Resampling.LANCZOS)
        destination.parent.mkdir(parents=True, exist_ok=True)
        image.save(destination, format="PNG")
    return {
        "source_size_px": {"width": source_size[0], "height": source_size[1]},
        "scaled_size_px": {"width": width_px, "height": height_px},
        "path": destination.name,
        "source_blob_sha256": source_blob_sha256,
        "preview_sha256": sha256(destination),
    }


def build_preview(docx_path: Path, output_dir: Path, dpi: float) -> dict[str, Any]:
    document = Document(str(docx_path))
    image_dir = output_dir / "images"
    image_dir.mkdir(parents=True, exist_ok=True)
    blocks: list[dict[str, Any]] = []
    images: list[dict[str, Any]] = []
    findings: list[dict[str, Any]] = []
    coverage = {
        "body_blocks": "walked_in_order",
        "nested_tables": "not_walked",
        "table_cell_images": "not_walked",
        "headers_footers": "not_walked",
        "floating_drawings": "not_walked",
        "word_pagination_fonts": "out_of_scope",
    }
    # The body walker intentionally handles top-level paragraphs/tables.  Do
    # not silently call that a complete final-document readback when OOXML
    # contains regions it cannot order or size safely.
    body = document.element.body
    if any(_local_name(node) == "anchor" for node in body.iter()):
        findings.append({"code": "floating_drawings_not_walked"})
    top_level_tables = [block for block in iter_block_items(document) if isinstance(block, Table)]
    for table in top_level_tables:
        nested = [node for node in table._tbl.iter() if _local_name(node) == "tbl"]
        if len(nested) > 1:
            findings.append({"code": "nested_tables_not_walked"})
            break
        if any(_local_name(node) == "blip" for node in table._tbl.iter()):
            findings.append({"code": "table_cell_images_not_walked"})
    for section in document.sections:
        for container_name, container in (("header", section.header), ("footer", section.footer)):
            if any(paragraph.text.strip() or paragraph._p.xpath(".//a:blip") for paragraph in container.paragraphs) or container.tables:
                findings.append({"code": "header_footer_not_walked", "region": container_name})
    if not any(True for _ in iter_block_items(document)):
        findings.append({"code": "content_empty"})
    image_index = 0
    for block_index, block in enumerate(iter_block_items(document)):
        if isinstance(block, Paragraph):
            entry: dict[str, Any] = {"index": block_index, "type": "paragraph", "text": block.text}
            occurrences = _image_occurrences(block, document)
            rendered_occurrences: list[dict[str, Any]] = []
            for occurrence in occurrences:
                image_index += 1
                image_id = f"image-{image_index:03d}"
                occurrence["id"] = image_id
                if occurrence.get("status") != "pass":
                    findings.append({"code": occurrence.get("reason", "image_unavailable"), "image": image_id})
                    rendered_occurrences.append(occurrence)
                    continue
                relation_id = occurrence["relationship"]
                part = document.part.related_parts[relation_id]
                extent = occurrence["inline_emu"]
                output_name = f"{image_id}.png"
                try:
                    scaled = _scale_image(
                        part,
                        image_dir / output_name,
                        (extent["cx"], extent["cy"]),
                        dpi,
                    )
                except (OSError, ValueError) as error:
                    occurrence.update(status="blocked", reason="image_decode_or_scale_failed")
                    findings.append({"code": "image_decode_or_scale_failed", "image": image_id})
                else:
                    occurrence.update(scaled)
                    occurrence["path"] = f"images/{output_name}"
                    images.append({
                        "id": image_id,
                        "path": occurrence["path"],
                        "part": occurrence["part"],
                        "inline_emu": occurrence["inline_emu"],
                        "source_size_px": occurrence["source_size_px"],
                        "scaled_size_px": occurrence["scaled_size_px"],
                        "source_blob_sha256": occurrence["source_blob_sha256"],
                        "preview_sha256": occurrence["preview_sha256"],
                        "sha256": sha256(image_dir / output_name),
                    })
                rendered_occurrences.append(occurrence)
            if rendered_occurrences:
                entry["images"] = rendered_occurrences
            blocks.append(entry)
        else:
            rows = [[cell.text for cell in row.cells] for row in block.rows]
            blocks.append({"index": block_index, "type": "table", "rows": rows})

    # Build a caption walk-through without removing the original ordered
    # paragraphs: an image paragraph followed by a text paragraph is the
    # builder's image/caption contract.
    for index, block in enumerate(blocks[:-1]):
        if block.get("type") != "paragraph" or not block.get("images"):
            continue
        following = blocks[index + 1]
        if following.get("type") == "paragraph" and following.get("text", "").strip():
            block["caption"] = following["text"]
            following["caption_for"] = [image["id"] for image in block["images"] if image.get("status") == "pass"]

    return {
        "schema_version": SCHEMA_VERSION,
        "docx": {"path": docx_path.name, "sha256": sha256(docx_path), "size_bytes": docx_path.stat().st_size},
        "dpi": dpi,
        "render_visual": {"status": "not_performed", "reason": "lite preview reads OOXML and inline images; it does not render Word pages."},
        "blocks": blocks,
        "images": images,
        "coverage": coverage,
        "findings": findings,
        "status": "blocked" if findings else "pass",
    }


def _html_preview(report: dict[str, Any]) -> str:
    body: list[str] = []
    for block in report["blocks"]:
        if block["type"] == "table":
            rows = []
            for row in block["rows"]:
                rows.append("<tr>" + "".join(f"<td>{html.escape(str(value))}</td>" for value in row) + "</tr>")
            body.append("<table>" + "".join(rows) + "</table>")
            continue
        text = block.get("text", "")
        if text:
            body.append(f"<p>{html.escape(text)}</p>")
        for image in block.get("images", []):
            if image.get("status") == "pass":
                body.append(f'<figure><img src="{html.escape(image["path"])}" alt="{html.escape(image["id"])}">')
                caption = block.get("caption")
                if caption:
                    body.append(f"<figcaption>{html.escape(caption)}</figcaption>")
                body.append("</figure>")
    title = html.escape(str(report["docx"]["path"]))
    return "<!doctype html><meta charset=\"utf-8\"><title>DOCX content preview</title>" \
        "<style>body{font-family:sans-serif;max-width:960px;margin:2rem auto}table{border-collapse:collapse;margin:1rem 0}td{border:1px solid #aaa;padding:.35rem}figure{margin:1rem 0}img{display:block;max-width:100%;height:auto}figcaption{font-style:italic}</style>" \
        f"<h1>{title}</h1><p>內容預覽，非Word分頁/字型渲染。</p><p>render_visual: not_performed</p>" + "".join(body)


def main(argv: list[str] | None = None) -> int:
    args = parse_args(argv)
    docx_path = args.docx.expanduser().resolve()
    output_dir = args.output_dir.expanduser().resolve()
    try:
        if not docx_path.is_file():
            raise FileNotFoundError(f"DOCX not found: {docx_path}")
        if args.dpi <= 0 or not math.isfinite(args.dpi) or args.dpi > MAX_DPI:
            raise ValueError(f"dpi must be a positive finite number no greater than {MAX_DPI:g}")
        if output_dir == docx_path or output_dir.is_relative_to(docx_path):
            raise ValueError("output directory must not be the DOCX file")
        if output_dir.exists() and any(output_dir.iterdir()):
            raise ValueError("output directory must be new or empty; refusing to overwrite preview/checkpoint artifacts")
        with zipfile.ZipFile(docx_path) as archive:
            archive.testzip()
        output_dir.mkdir(parents=True, exist_ok=True)
        report = build_preview(docx_path, output_dir, args.dpi)
        (output_dir / "preview.json").write_text(
            json.dumps(report, ensure_ascii=False, indent=2, sort_keys=True) + "\n", encoding="utf-8"
        )
        (output_dir / "preview.html").write_text(_html_preview(report), encoding="utf-8")
        print(f"PREVIEW_JSON={output_dir / 'preview.json'}")
        print(f"PREVIEW_HTML={output_dir / 'preview.html'}")
        return 0 if report["status"] == "pass" else 1
    except (OSError, ValueError, zipfile.BadZipFile) as error:
        print(f"preview_docx.py: {error}", file=sys.stderr)
        return 2


if __name__ == "__main__":
    raise SystemExit(main())
